from django.test import TestCase
from django.contrib.auth.models import User
from parking.models import Area, ParkingLot, ParkingUser, ParkingPrice, ParkingRegistrationRequest
from parking.utils.excel import (
    export_parkinglots_xlsx,
    export_parkingusers_xlsx,
    export_revenue_xlsx,
    export_registrations_xlsx,
)
from parking.utils.excel import (
    get_parkinglots_template_xlsx,
    get_parkingusers_template_xlsx,
    import_parkinglots_xlsx,
    import_parkingusers_xlsx,
)
import openpyxl
from io import BytesIO


class ExcelExportTest(TestCase):
    def setUp(self):
        self.area = Area.objects.create(
            name="Khu A", description="", latitude=10.0, longitude=106.0
        )
        self.lot = ParkingLot.objects.create(
            name="Bãi Test", address="123 Test", area=self.area,
            capacity=50, is_active=True, is_deleted=False,
        )
        self.user_auth = User.objects.create_user("staff1", password="pass")
        self.parking_user = ParkingUser.objects.create(
            full_name="Nguyễn Văn A", phone="0901234567",
            license_plate="51A-12345", vehicle_type="car",
            parking_lot=self.lot, is_active=True, is_deleted=False,
        )
        self.registration = ParkingRegistrationRequest.objects.create(
            full_name="Trần Thị B", phone="0912345678", email="b@test.com",
            address="456 Test", license_plate="59B-99999",
            vehicle_type="motorbike", parking_lot=self.lot,
            status="pending", created_by=self.user_auth,
        )

    def _load_wb(self, content: bytes):
        return openpyxl.load_workbook(BytesIO(content))

    def test_export_parkinglots_returns_bytes(self):
        result = export_parkinglots_xlsx()
        self.assertIsInstance(result, bytes)

    def test_export_parkinglots_has_header_row(self):
        wb = self._load_wb(export_parkinglots_xlsx())
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Tên bãi", headers)
        self.assertIn("Sức chứa", headers)

    def test_export_parkinglots_has_data_row(self):
        wb = self._load_wb(export_parkinglots_xlsx())
        ws = wb.active
        self.assertEqual(ws.max_row, 2)  # header + 1 lot

    def test_export_parkingusers_returns_bytes(self):
        result = export_parkingusers_xlsx()
        self.assertIsInstance(result, bytes)

    def test_export_parkingusers_has_header_row(self):
        wb = self._load_wb(export_parkingusers_xlsx())
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Họ tên", headers)
        self.assertIn("Biển số xe", headers)

    def test_export_parkingusers_has_data_row(self):
        wb = self._load_wb(export_parkingusers_xlsx())
        ws = wb.active
        self.assertEqual(ws.max_row, 2)  # header + 1 user

    def test_export_revenue_returns_bytes(self):
        result = export_revenue_xlsx()
        self.assertIsInstance(result, bytes)

    def test_export_revenue_has_header_row(self):
        wb = self._load_wb(export_revenue_xlsx())
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Tên bãi", headers)
        self.assertIn("Doanh thu ước tính (VNĐ/giờ)", headers)

    def test_export_revenue_has_data_row(self):
        wb = self._load_wb(export_revenue_xlsx())
        ws = wb.active
        self.assertEqual(ws.max_row, 2)  # header + 1 lot

    def test_export_registrations_returns_bytes(self):
        result = export_registrations_xlsx()
        self.assertIsInstance(result, bytes)

    def test_export_registrations_status_vietnamese(self):
        wb = self._load_wb(export_registrations_xlsx())
        ws = wb.active
        # Find status column index
        headers = [cell.value for cell in ws[1]]
        status_idx = headers.index("Trạng thái") + 1
        self.assertEqual(ws.cell(row=2, column=status_idx).value, "Chờ duyệt")


class ExcelTemplateTest(TestCase):
    def _load_wb(self, content: bytes):
        return openpyxl.load_workbook(BytesIO(content))

    def test_parkinglots_template_has_example_row(self):
        wb = self._load_wb(get_parkinglots_template_xlsx())
        ws = wb.active
        self.assertEqual(ws.max_row, 2)

    def test_parkingusers_template_has_example_row(self):
        wb = self._load_wb(get_parkingusers_template_xlsx())
        ws = wb.active
        self.assertEqual(ws.max_row, 2)


class ExcelImportTest(TestCase):
    def setUp(self):
        self.area = Area.objects.create(
            name="Khu Test", description="", latitude=10.0, longitude=106.0
        )
        self.lot = ParkingLot.objects.create(
            name="Bãi Hiện Có", address="123", area=self.area,
            capacity=10, is_active=True, is_deleted=False,
        )
        self.existing_user = ParkingUser.objects.create(
            full_name="User Cũ", phone="0900000000",
            license_plate="51A-00000", vehicle_type="car",
            parking_lot=self.lot, is_active=True, is_deleted=False,
        )

    def _make_parkinglots_xlsx(self, rows: list) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Tên bãi", "Địa chỉ", "Khu vực", "Quận/Huyện",
                   "Sức chứa", "Vĩ độ", "Kinh độ", "Mô tả ngắn", "Trạng thái"])
        for row in rows:
            ws.append(row)
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _make_parkingusers_xlsx(self, rows: list) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Họ tên", "Số điện thoại", "Email", "Địa chỉ",
                   "Biển số xe", "Loại xe", "Bãi đỗ"])
        for row in rows:
            ws.append(row)
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # --- ParkingLot import ---
    def test_import_parkinglots_success(self):
        data = self._make_parkinglots_xlsx([
            ["Bãi Mới", "456 Đường B", "Khu Test", "Q.1", 30, 10.1, 106.1, "Mô tả", 1]
        ])
        errors = import_parkinglots_xlsx(data)
        self.assertEqual(errors, [])
        self.assertTrue(ParkingLot.objects.filter(name="Bãi Mới").exists())

    def test_import_parkinglots_invalid_area(self):
        data = self._make_parkinglots_xlsx([
            ["Bãi X", "789", "Khu Không Tồn Tại", "", 20, None, None, "", 1]
        ])
        errors = import_parkinglots_xlsx(data)
        self.assertGreater(len(errors), 0)
        self.assertFalse(ParkingLot.objects.filter(name="Bãi X").exists())

    def test_import_parkinglots_invalid_capacity(self):
        data = self._make_parkinglots_xlsx([
            ["Bãi Y", "789", "Khu Test", "", "abc", None, None, "", 1]
        ])
        errors = import_parkinglots_xlsx(data)
        self.assertGreater(len(errors), 0)
        self.assertFalse(ParkingLot.objects.filter(name="Bãi Y").exists())

    # --- ParkingUser import ---
    def test_import_parkingusers_success(self):
        data = self._make_parkingusers_xlsx([
            ["Nguyễn Mới", "0988888888", "new@test.com", "Địa chỉ mới",
             "51C-99999", "car", "Bãi Hiện Có"]
        ])
        errors = import_parkingusers_xlsx(data)
        self.assertEqual(errors, [])
        self.assertTrue(ParkingUser.objects.filter(phone="0988888888").exists())

    def test_import_parkingusers_duplicate_phone(self):
        data = self._make_parkingusers_xlsx([
            ["Người Trùng", "0900000000", "", "", "51C-11111", "car", "Bãi Hiện Có"]
        ])
        errors = import_parkingusers_xlsx(data)
        self.assertGreater(len(errors), 0)
        self.assertEqual(ParkingUser.objects.filter(phone="0900000000").count(), 1)

    def test_import_parkingusers_duplicate_plate(self):
        data = self._make_parkingusers_xlsx([
            ["Người Trùng", "0977777777", "", "", "51A-00000", "car", "Bãi Hiện Có"]
        ])
        errors = import_parkingusers_xlsx(data)
        self.assertGreater(len(errors), 0)

    def test_import_parkingusers_invalid_vehicle_type(self):
        data = self._make_parkingusers_xlsx([
            ["Người Mới", "0966666666", "", "", "51D-22222", "truck", "Bãi Hiện Có"]
        ])
        errors = import_parkingusers_xlsx(data)
        self.assertGreater(len(errors), 0)

    def test_import_parkingusers_invalid_lot(self):
        data = self._make_parkingusers_xlsx([
            ["Người Mới", "0955555555", "", "", "51E-33333", "car", "Bãi Không Tồn Tại"]
        ])
        errors = import_parkingusers_xlsx(data)
        self.assertGreater(len(errors), 0)
